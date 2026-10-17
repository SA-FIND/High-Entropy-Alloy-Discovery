"""
Build the submission Word file and PDF from paper/manuscript_v2.md.

Figures are embedded after the paragraph that first cites them, with their captions; the
caption list stays at the end as Elsevier asks. Line numbers and page numbers are added.

Run (Windows Python with python-docx and pywin32, Word installed):
    python paper/build_docx.py [source.md output.docx]
Default: paper/manuscript_v2.md -> project_files/Solomon_Ahedor_MetaForge_Manuscript_v2.docx
(gitignored), plus the PDF. Other sources (e.g. the research proposal) need not cite every figure.
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "paper", "manuscript_v2.md")
OUT = os.path.join(ROOT, "project_files", "Solomon_Ahedor_MetaForge_Manuscript_v2.docx")
FIGS = {1: "fig1_dataset.png", 2: "fig2_chgnet_benchmark.png", 3: "fig3_parity_grouped.png",
        4: "fig4_model_comparison.png", 5: "fig5_residual_vs_T.png", 6: "fig6_mlip_elements.png",
        7: "fig7_pareto.png"}
WIDTH = {1: 15, 3: 16, 4: 15, 6: 15}  # cm; others 10


def add_runs(p, text):
    """Plain text with **bold** spans."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            p.add_run(part).bold = bool(i % 2)


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._r.append(el)


def line_numbers(section):
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
    section._sectPr.append(ln)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows if not re.match(r"^\|[-| ]+\|$", r.strip())]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, text in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, text)
            for r in p.runs:
                r.font.size = Pt(9)
                r.bold = r.bold or i == 0
    doc.add_paragraph()


def main():
    src, out = (os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])) if len(sys.argv) == 3 else (SRC, OUT)
    lines = open(src, encoding="utf-8").read().splitlines()
    captions = {}
    for ln in lines:
        m = re.match(r"^Fig\. (\d)\. (.*)", ln)
        if m:
            captions[int(m.group(1))] = m.group(2)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(11)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    page_number_footer(sec)
    if src == SRC:  # reviewers need line numbers; other documents do not
        line_numbers(sec)

    placed = set()
    in_captions = False
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("<!--"):
            while "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if not ln.strip():
            i += 1
            continue
        if ln.startswith("# "):
            p = doc.add_paragraph(); r = p.add_run(ln[2:]); r.bold = True; r.font.size = Pt(15)
        elif ln.startswith("## "):
            in_captions = ln[3:] == "Figure captions"
            if ln[3:] in ("Highlights", "Figure captions", "References"):
                doc.add_page_break()
            doc.add_heading(ln[3:], level=1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:], level=2)
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i]); i += 1
            table(doc, block)
            continue
        elif ln.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), ln[2:])
        elif re.match(r"^\d\. ", ln):
            # explicit numbers: Word's List Number style would continue numbering across lists
            p = doc.add_paragraph(); add_runs(p, ln)
            p.paragraph_format.left_indent = Cm(0.63); p.paragraph_format.first_line_indent = Cm(-0.63)
        elif re.match(r"^\[\d+\] ", ln):
            p = doc.add_paragraph(); add_runs(p, ln)
            p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(10)
        elif re.match(r"^Table \d\.", ln):
            p = doc.add_paragraph(); r = p.add_run(ln); r.bold = True; r.font.size = Pt(10)
        else:
            p = doc.add_paragraph(); add_runs(p, ln)
            if not in_captions:
                for n in sorted({int(x) for x in re.findall(r"Fig\. (\d)", ln)} - placed):
                    path = os.path.join(ROOT, "figures", FIGS[n])
                    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fp.add_run().add_picture(path, width=Cm(WIDTH.get(n, 10)))
                    cp = doc.add_paragraph(); r = cp.add_run(f"Fig. {n}. {captions[n]}"); r.font.size = Pt(9)
                    placed.add(n)
        i += 1

    if src == SRC:
        assert placed == set(FIGS), f"figures never cited: {set(FIGS) - placed}"
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)

    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(out)
        d.SaveAs2(out.replace(".docx", ".pdf"), FileFormat=17)
        pages = d.ComputeStatistics(2)
        d.Close(False)
    finally:
        word.Quit()
    print(f"wrote {out} and PDF, {pages} pages")


if __name__ == "__main__":
    main()
